"""平台保存时生成软件事件和硬件原始触发的独立工作表。"""
import os

import numpy as np


def hardware_runs(values, timestamps, counters, sample_rate):
    """记录每个连续非零原始值段；不阈值化、不与软件事件按序配对。

    缺失值会截断当前段，恢复后的非零值注明 after_missing，不推断缺失
    区间内是否发生触发。录制首样本已非零也不声称观测到了上升沿。
    """
    values = np.asarray(values, dtype=float)
    valid = np.isfinite(values)
    if not len(values):
        return []
    start = valid & (values != 0)
    start[1:] &= (~valid[:-1]) | (values[1:] != values[:-1])
    changes = np.r_[0, np.flatnonzero((valid[1:] != valid[:-1]) |
                                    (valid[1:] & valid[:-1] & (values[1:] != values[:-1]))) + 1,
                    len(values)]
    rows = []
    for first, end in zip(changes[:-1], changes[1:]):
        if not start[first]:
            continue
        origin = "initial_nonzero" if first == 0 else (
            "after_missing" if not valid[first - 1] else "value_change")
        rows.append(dict(hardware_event_id=len(rows) + 1, raw_value=float(values[first]),
                         sample_index=int(first), onset_sec=float(first / sample_rate),
                         timestamp_ms=float(timestamps[first]), sample_counter=int(counters[first]),
                         end_sample_index_exclusive=int(end), samples=int(end-first),
                         duration_sec=float((end-first) / sample_rate), origin=origin))
    return rows


def save_annotation_workbook(path, software_events, hardware_events, sample_rate, has_hardware):
    # This is a runtime export feature of the Python acquisition application.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    workbook = Workbook()
    workbook.remove(workbook.active)
    software = workbook.create_sheet("软件事件")
    hardware = workbook.create_sheet("硬件触发")
    sheets = [
        (software, software_events,
         ["event_id", "label", "event_type", "channel", "status", "sample_index", "onset_sec",
          "sample_time_ms", "timestamp_ms", "sample_counter", "source"],
         ["事件序号", "软件标签", "事件类型", "刺激通道", "定位状态", "采样索引(从0开始)", "相对时间(s)",
          "对应样本时间(ms)", "软件接收时间(ms)", "采样计数器", "来源"]),
        (hardware, hardware_events,
         ["hardware_event_id", "raw_value", "sample_index", "onset_sec", "timestamp_ms",
          "sample_counter", "end_sample_index_exclusive", "samples", "duration_sec", "origin"],
         ["非零段序号", "硬件原始值", "起始采样索引(从0开始)", "相对时间(s)", "样本时间(ms)",
          "采样计数器", "结束采样索引(不含)", "样本数", "持续时间(s)", "起点观测状态"]),
    ]
    for sheet, events, fields, headers in sheets:
        sheet.append(headers)
        for event in events:
            sheet.append([event.get(field) for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[cell.column_letter].width = 23
        for col, field in enumerate(fields, 1):
            fmt = "0.000000" if field.endswith("_sec") else (
                "0.000" if field.endswith("_ms") else "0" if field in {
                    "sample_index", "sample_counter", "end_sample_index_exclusive", "samples"} else None)
            if fmt:
                for cells in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                    cells[0].number_format = fmt
    software.cell(1, 12, "采样率(Hz)")
    software.cell(2, 12, float(sample_rate))
    hardware.cell(1, 11, "硬件通道存在")
    hardware.cell(2, 11, bool(has_hardware))
    hardware.cell(1, 12, "完整原始值见同名CSV的hardware_trigger列；本表仅列连续非零段")
    hardware.column_dimensions["L"].width = 65
    temporary = path + ".tmp.xlsx"
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.remove(temporary)
